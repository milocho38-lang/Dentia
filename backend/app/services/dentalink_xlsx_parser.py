from __future__ import annotations

import hashlib
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 48 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_ROWS = 5_000
MAX_COLUMNS = 50
MAX_CELL_LENGTH = 3_000
SOURCE_SYSTEM = "DENTALINK"


class DentalinkWorkbookError(ValueError):
    pass


VERIFIED_HEADERS = (
    "# Paciente",
    "# Interno",
    "RUT",
    "Nombre",
    "Apellidos",
    "Fecha de nac.",
    "Edad",
    "Teléfono",
    "Celular",
    "Ciudad",
    "Comuna",
    "Dirección",
    "E-Mail",
    "Alertas",
    "Observaciones",
    "Sexo",
    "Tipo Paciente",
    "# Apoderado",
    "Convenio",
    "Nombre Empresa Convenio",
    "Empleador",
    "Referencia",
    "Nacionalidad",
    "Migrante",
    "Pueblos Originarios",
)


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", text.strip().casefold())
    plain = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return " ".join(plain.split())


HEADER_BY_NORMALIZED = {normalize_header(item): item for item in VERIFIED_HEADERS}
ESSENTIAL_HEADERS = {"# Paciente", "Nombre", "Apellidos"}
PENDING_HEADERS = {
    "# Interno": "internal_patient_number",
    "Alertas": "alerts",
    "Observaciones": "observations",
    "Tipo Paciente": "patient_type",
    "Convenio": "agreement",
    "Nombre Empresa Convenio": "agreement_company",
    "Empleador": "employer",
    "Referencia": "reference",
    "Nacionalidad": "nationality",
    "Migrante": "migrant_status",
    "Pueblos Originarios": "indigenous_identity",
    "Comuna": "commune",
    "# Apoderado": "responsible_source_reference",
}


@dataclass
class ParsedDentalinkRow:
    row_number: int
    source_patient_id: str
    first_names: str
    last_names: str
    rut: str | None
    birth_date: date | None
    mobile: str
    alternate_phone: str | None
    email: str | None
    city: str | None
    address: str | None
    sex: str | None
    status: str = "READY"
    issues: list[str] = field(default_factory=list)
    pending_fields: list[str] = field(default_factory=list)

    @property
    def display_name(self) -> str:
        return f"{self.first_names} {self.last_names}".strip()


@dataclass
class ParsedDentalinkWorkbook:
    file_sha256: str
    sheet_name: str
    rows: list[ParsedDentalinkRow]
    extra_headers: list[str]
    missing_headers: list[str]


def _cell_text(value: object, *, max_length: int = MAX_CELL_LENGTH) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) > max_length:
        raise DentalinkWorkbookError("Una celda supera el tamaño permitido.")
    return text


def _source_id(value: object) -> str:
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _cell_text(value, max_length=120)


def _inspect_archive(content: bytes) -> None:
    if len(content) > MAX_FILE_BYTES:
        raise DentalinkWorkbookError("El archivo supera 8 MB.")
    try:
        archive = zipfile.ZipFile(BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise DentalinkWorkbookError("El archivo no es un XLSX válido.") from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise DentalinkWorkbookError("El archivo contiene demasiadas entradas internas.")
        total_size = 0
        for item in entries:
            normalized = item.filename.replace("\\", "/")
            if normalized.startswith("/") or ".." in normalized.split("/"):
                raise DentalinkWorkbookError("El XLSX contiene rutas internas no válidas.")
            if item.flag_bits & 0x1:
                raise DentalinkWorkbookError("No se admiten archivos Excel cifrados.")
            total_size += item.file_size
            if total_size > MAX_UNCOMPRESSED_BYTES:
                raise DentalinkWorkbookError("El XLSX expandido supera el límite seguro.")
            if item.file_size > 12 * 1024 * 1024:
                raise DentalinkWorkbookError("Una entrada interna supera el límite seguro.")
            lowered = normalized.casefold()
            if (
                "vbaproject" in lowered
                or lowered.startswith("xl/externallinks/")
                or lowered.startswith("xl/embeddings/")
                or "/oleobjects/" in lowered
                or lowered == "xl/connections.xml"
            ):
                raise DentalinkWorkbookError(
                    "No se admiten macros, enlaces externos ni objetos incrustados."
                )
            if lowered.startswith("xl/worksheets/") and lowered.endswith(".xml"):
                raw = archive.read(item)
                if re.search(br"<(?:[A-Za-z0-9_]+:)?f(?:\s|>)", raw):
                    raise DentalinkWorkbookError("No se admiten fórmulas en el Excel.")
            if lowered.endswith(".rels"):
                raw = archive.read(item)
                if re.search(br"TargetMode\s*=\s*['\"]External['\"]", raw, re.IGNORECASE):
                    raise DentalinkWorkbookError("No se admiten enlaces externos en el Excel.")
        if archive.testzip() is not None:
            raise DentalinkWorkbookError("El XLSX está dañado.")


def normalize_rut(value: str) -> str | None:
    compact = re.sub(r"[^0-9Kk]", "", value).upper()
    return compact or None


def is_valid_rut(value: str) -> bool:
    compact = normalize_rut(value)
    if compact is None or not re.fullmatch(r"\d{7,8}[0-9K]", compact):
        return False
    body, verifier = compact[:-1], compact[-1]
    total = 0
    factor = 2
    for digit in reversed(body):
        total += int(digit) * factor
        factor = 2 if factor == 7 else factor + 1
    result = 11 - (total % 11)
    expected = "0" if result == 11 else "K" if result == 10 else str(result)
    return verifier == expected


def _parse_date(value: object) -> tuple[date | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = _cell_text(value, max_length=40)
        parsed = None
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        if parsed is None:
            return None, "Fecha de nacimiento no interpretable."
    if parsed > date.today():
        return None, "Fecha de nacimiento futura."
    if parsed.year < 1900:
        return None, "Fecha de nacimiento fuera del rango admitido."
    return parsed, None


def _parse_email(value: object) -> tuple[str | None, str | None]:
    text = _cell_text(value, max_length=200)
    if not text:
        return None, None
    normalized = text.casefold()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized):
        return None, "Correo inválido; se omitirá."
    return normalized, None


def _parse_sex(value: object) -> tuple[str | None, str | None]:
    normalized = normalize_header(value)
    if not normalized or normalized in {"no informa", "sin informacion", "n/a"}:
        return None, None
    mapping = {
        "f": "femenino",
        "femenino": "femenino",
        "mujer": "femenino",
        "m": "masculino",
        "masculino": "masculino",
        "hombre": "masculino",
        "otro": "otro",
    }
    if normalized not in mapping:
        return None, "Sexo no reconocido; se omitirá."
    return mapping[normalized], None


def _mark(row: ParsedDentalinkRow, status: str, issue: str) -> None:
    priority = {"READY": 0, "INCOMPLETE": 1, "REVIEW": 2, "REJECTED": 3}
    if priority[status] > priority[row.status]:
        row.status = status
    if issue not in row.issues:
        row.issues.append(issue)


def parse_dentalink_workbook(content: bytes) -> ParsedDentalinkWorkbook:
    _inspect_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise DentalinkWorkbookError("No fue posible leer el XLSX.") from exc
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise DentalinkWorkbookError("El archivo debe contener la hoja Sheet1.")
        worksheet = workbook["Sheet1"]
        if worksheet.max_column > MAX_COLUMNS or worksheet.max_row > MAX_ROWS + 1:
            raise DentalinkWorkbookError("El archivo supera el límite de filas o columnas.")
        header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        normalized_seen: dict[str, int] = {}
        column_by_header: dict[str, int] = {}
        extra_headers: list[str] = []
        for index, raw_header in enumerate(header_values):
            normalized = normalize_header(raw_header)
            if not normalized:
                continue
            if normalized in normalized_seen:
                raise DentalinkWorkbookError("El archivo contiene encabezados duplicados o ambiguos.")
            normalized_seen[normalized] = index
            canonical = HEADER_BY_NORMALIZED.get(normalized)
            if canonical is None:
                extra_headers.append(_cell_text(raw_header, max_length=120))
            else:
                column_by_header[canonical] = index
        missing_headers = [item for item in VERIFIED_HEADERS if item not in column_by_header]
        missing_essential = sorted(ESSENTIAL_HEADERS.intersection(missing_headers))
        if missing_essential:
            raise DentalinkWorkbookError(
                "Faltan encabezados obligatorios: " + ", ".join(missing_essential)
            )

        def value(values: tuple, header: str) -> object:
            index = column_by_header.get(header)
            return values[index] if index is not None and index < len(values) else None

        parsed_rows: list[ParsedDentalinkRow] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if row_number > MAX_ROWS + 1 or len(values) > MAX_COLUMNS:
                raise DentalinkWorkbookError("El archivo supera el límite de filas o columnas.")
            if not any(item not in (None, "") for item in values):
                continue
            source_id = _source_id(value(values, "# Paciente"))
            first_names = _cell_text(value(values, "Nombre"), max_length=150)
            last_names = _cell_text(value(values, "Apellidos"), max_length=150)
            birth_date, date_issue = _parse_date(value(values, "Fecha de nac."))
            email, email_issue = _parse_email(value(values, "E-Mail"))
            sex, sex_issue = _parse_sex(value(values, "Sexo"))
            rut_text = _cell_text(value(values, "RUT"), max_length=50)
            row = ParsedDentalinkRow(
                row_number=row_number,
                source_patient_id=source_id,
                first_names=first_names,
                last_names=last_names,
                rut=normalize_rut(rut_text),
                birth_date=birth_date,
                mobile=_cell_text(value(values, "Celular"), max_length=50),
                alternate_phone=_cell_text(value(values, "Teléfono"), max_length=50) or None,
                email=email,
                city=_cell_text(value(values, "Ciudad"), max_length=100) or None,
                address=_cell_text(value(values, "Dirección"), max_length=300) or None,
                sex=sex,
            )
            if not source_id or not first_names or not last_names:
                _mark(row, "REJECTED", "Falta # Paciente, Nombre o Apellidos.")
            if rut_text and not is_valid_rut(rut_text):
                row.rut = None
                _mark(row, "REVIEW", "RUT inválido; requiere revisión.")
            if not row.rut:
                _mark(row, "INCOMPLETE", "Paciente sin RUT.")
            if not row.mobile:
                _mark(row, "INCOMPLETE", "Paciente sin celular.")
            if row.birth_date is None:
                _mark(row, "INCOMPLETE", date_issue or "Paciente sin fecha de nacimiento.")
            if email_issue:
                _mark(row, "INCOMPLETE", email_issue)
                row.pending_fields.append("invalid_email")
            if sex_issue:
                _mark(row, "INCOMPLETE", sex_issue)
                row.pending_fields.append("unrecognized_sex")
            for header, pending_name in PENDING_HEADERS.items():
                if _cell_text(value(values, header)):
                    row.pending_fields.append(pending_name)
            if _cell_text(value(values, "Edad")) and row.birth_date is None:
                row.pending_fields.append("age_without_birth_date")
            if row.birth_date is not None:
                today = date.today()
                age = today.year - row.birth_date.year - (
                    (today.month, today.day) < (row.birth_date.month, row.birth_date.day)
                )
                if age < 18:
                    _mark(row, "REVIEW", "Menor sin responsable completo importable.")
            parsed_rows.append(row)

        source_counts = Counter(row.source_patient_id for row in parsed_rows if row.source_patient_id)
        rut_counts = Counter(row.rut for row in parsed_rows if row.rut)
        date_counts = Counter(row.birth_date for row in parsed_rows if row.birth_date)
        placeholder_dates = {
            item
            for item, count in date_counts.items()
            if count >= 20 and count / max(len(parsed_rows), 1) >= 0.05
        }
        for row in parsed_rows:
            if row.source_patient_id and source_counts[row.source_patient_id] > 1:
                _mark(row, "REJECTED", "# Paciente duplicado dentro del archivo.")
            if row.rut and rut_counts[row.rut] > 1:
                _mark(row, "REVIEW", "RUT duplicado dentro del archivo.")
            if row.birth_date in placeholder_dates:
                row.birth_date = None
                _mark(row, "REVIEW", "Fecha repetida masivamente; posible placeholder.")
        return ParsedDentalinkWorkbook(
            file_sha256=hashlib.sha256(content).hexdigest(),
            sheet_name=worksheet.title,
            rows=parsed_rows,
            extra_headers=extra_headers,
            missing_headers=missing_headers,
        )
    finally:
        workbook.close()
